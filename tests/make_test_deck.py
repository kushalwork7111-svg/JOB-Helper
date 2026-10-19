"""Build a synthetic messy deck for testing the add-in: tests/decks/messy.pptx.

No client content: every slide is made-up text with planted formatting problems.
Run: python3 tests/make_test_deck.py   (needs: pip install python-pptx)
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

OUT = Path(__file__).parent / "decks" / "messy.pptx"


def box(slide, left, top, width, height, text, size, font="Arial", color=None, bold=False):
    tb = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(width), Pt(height))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return tb


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]

    # 1. Clean reference slide: title placeholder in the standard spot.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Executive summary"
    box(s, 36, 100, 888, 300, "Sample body text at the brand size.", 14)

    # 2. Title text box off position, wrong font and size, ALL CAPS.
    s = prs.slides.add_slide(blank)
    box(s, 60, 50, 700, 60, "MARKET OVERVIEW FOR FY26", 30, font="Calibri", bold=True)
    box(s, 36, 130, 888, 200, "Body in Calibri at 16 pt.", 16, font="Calibri")

    # 3. Off-palette colour, tiny text, mixed body sizes.
    s = prs.slides.add_slide(blank)
    box(s, 36, 28, 888, 52, "Scope of work", 24, bold=True)
    box(s, 36, 100, 420, 200, "Body at 12 pt in an off-brand blue.", 12, color="1E90FF")
    box(s, 480, 100, 420, 200, "Body at 18 pt.", 18)
    box(s, 36, 330, 420, 40, "Tiny 8 pt note in the middle of the slide.", 8)

    # 4. Empty text box, a box running off the slide, and a source line.
    s = prs.slides.add_slide(blank)
    box(s, 36, 28, 888, 52, "Approach", 24, bold=True)
    box(s, 36, 100, 300, 100, "", 14)
    box(s, 800, 200, 300, 100, "This box runs off the right edge.", 14)
    box(s, 36, 505, 600, 20, "Source: synthetic test data", 9)

    # 5. No title at all (picture-style slide with only bottom text).
    s = prs.slides.add_slide(blank)
    box(s, 36, 300, 888, 60, "Only body text, low on the slide.", 14)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
